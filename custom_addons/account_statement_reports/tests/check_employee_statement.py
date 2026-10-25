"""Execute report logic on ORM-shaped fixtures; not an installed Odoo test."""
import sys,types,importlib.util,unittest,base64
from pathlib import Path
from datetime import date
from io import BytesIO
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
class UserError(Exception):pass
class ValidationError(Exception):pass
class Dates:
    def __new__(cls,*a,**k):return None
    @staticmethod
    def to_date(v):return date.fromisoformat(v) if isinstance(v,str) else v
class Fields:
    Date=Dates
    def __getattr__(self,n):return lambda *a,**k:None
odoo=types.ModuleType('odoo');odoo.models=types.SimpleNamespace(Model=object,AbstractModel=object);odoo.fields=Fields()
odoo.api=types.SimpleNamespace(model=lambda f:f,onchange=lambda *a:lambda f:f,constrains=lambda *a:lambda f:f)
exceptions=types.ModuleType('odoo.exceptions');exceptions.UserError=UserError;exceptions.ValidationError=ValidationError
sys.modules.update({'odoo':odoo,'odoo.exceptions':exceptions})
package=types.ModuleType('asr_fixture');package.__path__=[str(ROOT/'models')];sys.modules[package.__name__]=package
def load(name):
 spec=importlib.util.spec_from_file_location('asr_fixture.'+name,ROOT/'models'/(name+'.py'));m=importlib.util.module_from_spec(spec);sys.modules[spec.name]=m;spec.loader.exec_module(m);return m
mixin_module=load('statement_mixin');statement_module=load('account_statement')
class Records(list):
    @property
    def ids(self):return [r.id for r in self]
    def filtered(self,f):return Records(r for r in self if f(r))
    def mapped(self,k):return [getattr(r,k) for r in self]
    def unlink(self):self.clear()
def rec(**kw):return types.SimpleNamespace(**kw)
def match(row,domain):
 for key,op,value in domain:
  actual=row
  for name in key.split('.'):actual=getattr(actual,name)
  if hasattr(actual,'id'):actual=actual.id
  if op=='=' and actual!=value:return False
  if op=='in' and actual not in value:return False
  if op=='<' and actual>=Dates.to_date(value):return False
  if op=='>=' and actual<Dates.to_date(value):return False
  if op=='<=' and actual>Dates.to_date(value):return False
 return True
class MoveLines:
 def __init__(self,rows):self.rows=rows
 def search(self,domain,order=None):return sorted([r for r in self.rows if match(r,domain)],key=lambda r:(r.date,r.id))
class Statements(unittest.TestCase):
 def setUp(self):
  self.company=rec(id=1,display_name='ABSAR')
  self.partner=rec(id=10,name='Employee',display_name='Employee')
  self.benefits=rec(id=100,account_type='liability_current',display_name='Accrued Benefits',company_ids=[self.company])
  self.advance=rec(id=101,account_type='asset_current',display_name='Employee Advances',company_ids=[self.company])
  self.expense=rec(id=102,account_type='expense',display_name='Benefit Expense',company_ids=[self.company])
  self.ar=rec(id=103,account_type='asset_receivable',display_name='AR',company_ids=[self.company])
  self.ap=rec(id=104,account_type='liability_payable',display_name='AP',company_ids=[self.company])
  def row(i,account,debit=0,credit=0,day='2026-10-01',partner=10,company=1,state='posted'):
   return rec(id=i,account_id=account,debit=debit,credit=credit,date=date.fromisoformat(day),partner_id=partner,company_id=company,parent_state=state,journal_id=1,name='Annual benefit' if account==self.benefits else 'Other',date_maturity=None,move_id=rec(id=i,name='JV/%s'%i,move_type='entry',ref='Ref',payment_reference=''))
  self.rows=[row(1,self.benefits,credit=1000,day='2026-09-01'),row(2,self.benefits,credit=5000),row(3,self.expense,debit=5000),row(4,self.advance,debit=2000),row(5,self.benefits,debit=500),row(6,self.benefits,credit=999,partner=11),row(7,self.benefits,credit=888,company=2),row(8,self.benefits,credit=100,state='draft'),row(9,self.ar,debit=50),row(10,self.ap,credit=75),row(11,self.benefits,credit=777,state='cancel')]
  self.mixin=mixin_module.StatementMixin();self.mixin.env={'account.move.line':MoveLines(self.rows)}
 def lines(self,kind='employee',accounts=None,state='posted',opening=True):
  return self.mixin._get_statement_lines_with_balance(self.partner,kind,'2026-10-01','2026-10-31',self.company,account_ids=accounts or [100,101,102],move_state=state,show_opening_balance=opening)
 def test_employee_opening_and_balance(self):
  lines=self.lines();self.assertEqual([x['move_id'] for x in lines],[False,2,4,5]);self.assertEqual(lines[0]['balance'],-1000);self.assertEqual(lines[-1]['balance'],-3500)
  self.assertEqual(lines[1]['account_id'],100);self.assertEqual(lines[1]['description'],'Annual benefit')
 def test_no_double_counting_and_scope(self):
  ids=[x['move_id'] for x in self.lines()];self.assertNotIn(3,ids);self.assertNotIn(6,ids);self.assertNotIn(7,ids);self.assertNotIn(8,ids);self.assertNotIn(11,ids)
  self.assertEqual([x['move_id'] for x in self.lines(accounts=[101])],[False,4])
 def test_draft_option_excludes_cancelled(self):
  lines=self.lines(state='all');self.assertIn(8,[x['move_id'] for x in lines]);self.assertNotIn(11,[x['move_id'] for x in lines]);self.assertEqual(lines[-1]['balance'],-3600)
 def test_hidden_opening_still_carried(self):
  lines=self.lines(opening=False);self.assertEqual(len(lines),3);self.assertEqual(lines[-1]['balance'],-3500)
 def test_customer_vendor_regression(self):
  for typ,ident,amount in [('asset_receivable',9,50),('liability_payable',10,-75)]:
   lines=self.mixin._get_statement_lines_with_balance(self.partner,typ,'2026-10-01','2026-10-31',self.company);self.assertEqual([x['move_id'] for x in lines],[False,ident]);self.assertEqual(lines[-1]['balance'],amount)
 def test_retention_in_customer_and_vendor_statements(self):
  for typ,accounts,expected,opening,closing in [
   ('asset_receivable',[103,101],[False,4,9],0,2050),
   ('liability_payable',[104,100],[False,2,5,10],-1000,-5575),
  ]:
   lines=self.lines(kind=typ,accounts=accounts)
   self.assertEqual([line['move_id'] for line in lines],expected)
   self.assertEqual(lines[0]['balance'],opening)
   self.assertEqual(lines[-1]['balance'],closing)
  # An explicit retention-only selection must not silently add ordinary AR/AP.
  self.assertEqual([line['move_id'] for line in self.lines(kind='asset_receivable',accounts=[101])],[False,4])
  # Even crafted selections cannot add expense counterparts to the balance.
  self.assertEqual([line['move_id'] for line in self.lines(kind='asset_receivable',accounts=[101,102])],[False,4])
 def test_retention_apply_and_validation(self):
  for typ,accounts,total in [('receivable',[self.ar,self.advance],2050),('payable',[self.ap,self.benefits],-5575)]:
   s=self.statement(typ,accounts);s.action_apply_filters();self.assertEqual(s.final_balance,total)
   for bad in [[self.expense],[rec(id=999,account_type='asset_current',company_ids=[rec(id=2)])]]:
    with self.assertRaises(UserError):self.statement(typ,bad).action_apply_filters()
 def test_retention_account_picker_domains(self):
  from lxml import etree
  tree=etree.parse(str(ROOT/'views/statement_views.xml'))
  field=tree.xpath('//field[@name="account_ids"]')[0]
  domain=eval(field.get('domain'),{},dict(company_id=1))
  self.assertIn(('account_type','in',list(mixin_module.EMPLOYEE_ACCOUNT_TYPES)),domain)
  for kind in ('customer','vendor'):
   text=(ROOT/'models'/(kind+'_statement_wizard.py')).read_text()
   self.assertIn("('company_ids', 'in', [company_id])",text)
   self.assertIn("'asset_current'",text);self.assertIn("'liability_current'",text)
 def test_employee_requires_explicit_accounts(self):
  with self.assertRaises(UserError):self.mixin._base_domain(self.partner,'employee',self.company)
 def statement(self,typ='employee',accounts=None):
  s=statement_module.AccountStatement();s.ensure_one=lambda:None;s.statement_type=typ;s.partner_id=self.partner;s.company_id=self.company;s.account_ids=Records(accounts if accounts is not None else [self.benefits,self.advance]);s.journal_ids=Records();s.date_from=date(2026,10,1);s.date_to=date(2026,10,31);s.move_state='posted';s.show_opening_balance=True;s.id=1;s.line_ids=Records()
  class Lines:
   def create(inner,vals):s.line_ids.extend(rec(account_id=next((a for a in [self.benefits,self.advance,self.expense,self.ar,self.ap] if a.id==v.get('account_id')),False),**{k:v for k,v in v.items() if k!='account_id'}) for v in vals)
  s.env={'statement.mixin':self.mixin,'account.statement.line':Lines()};return s
 def test_apply_and_invalid_accounts(self):
  s=self.statement();s.action_apply_filters();self.assertEqual(s.line_count,3);self.assertEqual(s.final_balance,-3500);self.assertEqual(s.total_debit,2500);self.assertEqual(s.total_credit,6000)
  for accounts in [[],[self.expense],[rec(id=999,account_type='liability_current',company_ids=[rec(id=2)])]]:
   with self.assertRaises(UserError):self.statement(accounts=accounts).action_apply_filters()
 def test_switch_type_starts_fresh(self):
  s=self.statement();s.action_apply_filters();s._default_period=lambda:(date(2026,10,1),date(2026,10,24))
  s._clear_results=lambda: [setattr(s,k,0) for k in ('opening_balance','final_balance','total_debit','total_credit','line_count')]
  s._onchange_statement_type()
  self.assertFalse(s.partner_id);self.assertEqual(s.line_ids,[(5,0,0)]);self.assertEqual(s.account_ids,[(5,0,0)]);self.assertEqual(s.journal_ids,[(5,0,0)])
  self.assertEqual(s.final_balance,0);self.assertEqual(s.move_state,'posted');self.assertEqual(s.date_to,date(2026,10,24))
 def test_type_buttons_reset_only_when_changed(self):
  s=self.statement();writes=[];resets=[]
  s.write=lambda vals:(writes.append(vals),s.__dict__.update(vals))
  s.action_reset_filters=lambda:(resets.append(True) or False)
  s.action_type_payable();s.action_type_receivable();s.action_type_employee()
  self.assertEqual([v['statement_type'] for v in writes],['payable','receivable','employee']);self.assertEqual(len(resets),3)
  s.action_type_employee();self.assertEqual(len(resets),3)
 def test_excel_headers_and_totals(self):
  for typ,accounts,size,label in [('employee',[self.benefits,self.advance],8,'Employee'),('receivable',[self.ar],7,'Customer'),('payable',[self.ap],7,'Vendor'),('receivable',[self.ar,self.advance],7,'Customer'),('payable',[self.ap,self.benefits],7,'Vendor')]:
   s=self.statement(typ,accounts);s.action_apply_filters();attachments=[]
   s._fields={'statement_type':rec(selection=[('receivable','Customer'),('payable','Vendor'),('employee','Employee')])}
   s.env['ir.attachment']=rec(create=lambda vals:(attachments.append(vals) or rec(id=22)))
   s.action_export_excel();book=openpyxl.load_workbook(BytesIO(base64.b64decode(attachments[0]['datas'])));sheet=book.active
   self.assertEqual(sheet.cell(3,2).value,label);self.assertEqual(sheet.cell(10,size).value,'Balance');self.assertEqual(sheet.cell(11+len(s.line_ids),size).value,s.final_balance)
   if typ=='employee':self.assertEqual(sheet.cell(10,3).value,'Account');self.assertEqual(sheet.cell(12,4).value,'Ref')
if __name__=='__main__':unittest.main(verbosity=2)
